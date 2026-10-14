"""Regression tests for Brand Share Cons step 1 (Category Topline's Total/L0/L1/L2 numbers
propagate into Brand Share Cons' matching rows) and step 2 (brand rows within each L2/flat-L0
group get their own ETS+Trend+expert-alignment+reconciliation forecast, one level deeper).
"""
import hashlib
import json
import os

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

import app as appmod


def _row(row_num, kind, l0=None, l1=None, l2=None, label=None, share=None, gmv=None):
    return {"row_num": row_num, "kind": kind, "l0": l0, "l1": l1, "l2": l2,
            "label": label or l2 or l1 or l0 or "Total", "share": share or [], "gmv": gmv or []}


# ---------------------------------------------------------------- _compute_brand_share_forecast

def test_brand_group_without_others_row_is_left_untouched():
    bsc_rows = [
        _row(1, "l0", l0="Cat A"),
        _row(2, "brand", l0="Cat A", l1="Cat A", l2="Cat A", label="BrandX", share=[0.5] * 12, gmv=[50] * 12),
        _row(3, "brand", l0="Cat A", l1="Cat A", l2="Cat A", label="BrandY", share=[0.5] * 12, gmv=[50] * 12),
    ]
    category_forecast = {appmod._row_key(bsc_rows[0]): {"gmv": 200.0}}
    out = appmod._compute_brand_share_forecast(bsc_rows, category_forecast)
    assert out == {}  # no "others" row in this group -> nothing computed, same as the edit view's rule


def test_brand_group_with_others_row_reconciles_to_parent_gmv():
    bsc_rows = [
        _row(1, "l0", l0="Cat A"),
        _row(2, "brand", l0="Cat A", l1="Cat A", l2="Cat A", label="BrandX", share=[0.7] * 12, gmv=[70] * 12),
        _row(3, "others", l0="Cat A", l1="Cat A", l2="Cat A", label="Others", share=[0.3] * 12, gmv=[30] * 12),
    ]
    category_forecast = {appmod._row_key(bsc_rows[0]): {"gmv": 200.0}}
    out = appmod._compute_brand_share_forecast(bsc_rows, category_forecast)

    key_x, key_others = appmod._row_key(bsc_rows[1]), appmod._row_key(bsc_rows[2])
    assert key_x in out and key_others in out
    assert out[key_x]["share"] == pytest.approx(0.7, abs=1e-6)  # flat history -> stays at 0.7
    assert out[key_others]["is_plug"] is True
    assert out[key_others]["parent_row"] == 1  # the L0 row's own row_num
    assert out[key_others]["sibling_rows"] == [2]
    # GMV cascades from the group's parent (200, the just-computed category total).
    assert out[key_x]["gmv"] == pytest.approx(140.0, abs=1e-4)
    assert out[key_others]["gmv"] == pytest.approx(60.0, abs=1e-4)
    assert out[key_x]["gmv"] + out[key_others]["gmv"] == pytest.approx(200.0, abs=1e-6)


def test_brand_expert_alignment_respects_cap_and_floor():
    bsc_rows = [
        _row(1, "l0", l0="Cat A"),
        _row(2, "brand", l0="Cat A", l1="Cat A", l2="Cat A", label="BrandX", share=[0.6] * 12, gmv=[600.0] * 12),
        _row(3, "others", l0="Cat A", l1="Cat A", l2="Cat A", label="Others", share=[0.4] * 12, gmv=[400.0] * 12),
    ]
    category_forecast = {appmod._row_key(bsc_rows[0]): {"gmv": 1200.0}}
    key_x = appmod._row_key(bsc_rows[1])

    # Within the 3pp flag threshold (740 vs year-ago 600 = +23.3%, double-digit -> +/-2pp band,
    # clipping base growth 0.2 up to 0.213 -> raw candidate 0.6067, a 0.67pp move) -> applied, not
    # flagged.
    lookup_small_move = {("cat a", "cat a", "cat a", "brandx"): 740.0}
    out = appmod._compute_brand_share_forecast(bsc_rows, category_forecast, lookup_small_move)
    assert out[key_x]["expert_flag"] is False
    assert out[key_x]["share"] > 0.6

    # Past the 4pp hard cap but within the 8pp disagreement gate (792 vs 600 = +32% YoY, clips to
    # 0.30 -> raw candidate 0.65, a 5pp move; raw unclipped implied share is 0.66, a 6pp move,
    # still under the gate) -> raw candidate hard-capped to 0.6+0.04=0.64, flagged, giving BrandX a
    # target GMV of 0.64*1200=768 (Others' target is unchanged at 480; together 48 over the
    # category's 1200). The group is jointly reconciled, weighted 1/yearago^2 per brand (BrandX=
    # 600, Others=400) rather than one flat scale factor, so BrandX absorbs more of that gap than
    # Others: final = 768 - 48*(600^2/(600^2+400^2)) = 734.7692..., share = 734.7692.../1200 - same
    # bounded-reconciliation math as Category Topline's resolve_level, ported here.
    lookup_big_move = {("cat a", "cat a", "cat a", "brandx"): 792.0}
    out2 = appmod._compute_brand_share_forecast(bsc_rows, category_forecast, lookup_big_move)
    assert out2[key_x]["expert_flag"] is True
    assert out2[key_x]["share"] == pytest.approx(0.6123076923076922, abs=1e-6)


# ---------------------------------------------------------------- end-to-end via set-total-gmv

def _write_auth_file(path, password):
    salt = os.urandom(16).hex()
    h = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt), 200_000).hex()
    path.write_text(json.dumps({"salt": salt, "hash": h, "algo": "pbkdf2_sha256_200000"}), encoding="utf-8")


@pytest.fixture
def client(tmp_path, monkeypatch):
    portal_dir = tmp_path / "Portal Data"
    tagged_dir = portal_dir / "Tagged Files"
    tagged_dir.mkdir(parents=True)

    monkeypatch.setattr(appmod, "PORTAL_DATA_DIR", portal_dir)
    monkeypatch.setattr(appmod, "PORTAL_TAGGED_DIR", tagged_dir)
    monkeypatch.setattr(appmod, "MONTHLY_STATE_FILE", portal_dir / "monthly_platform_state.json")
    monkeypatch.setattr(appmod, "ADMIN_AUTH_FILE", portal_dir / "admin_auth.json")
    monkeypatch.setattr(appmod, "SL_AUTH_FILE", portal_dir / "sl_auth.json")
    monkeypatch.setattr(appmod, "MANAGER_AUTH_FILE", portal_dir / "manager_auth.json")
    _write_auth_file(appmod.ADMIN_AUTH_FILE, "admin-pw")
    _write_auth_file(appmod.SL_AUTH_FILE, "sl-pw")
    _write_auth_file(appmod.MANAGER_AUTH_FILE, "manager-pw")

    monkeypatch.setattr(appmod, "_workbook_platforms", lambda path: ["Blinkit"])

    topline_dir = portal_dir / "Category Topline"
    topline_dir.mkdir(parents=True)
    topline_all_path = topline_dir / "Category Topline - All Platforms.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Blinkit")
    months = ["GMV Oct-30", "GMV Nov-30"]
    ws.append(["Kind", "Level", "Label", "L0", "L1", "L2"] + months)
    ws.append(["total", 0, "Total BPC", None, None, None, 100, 110])
    ws.append(["l0", 1, "Cat A", "Cat A", None, None, 60, 66])
    ws.append(["l0_flat", 1, "Cat B", "Cat B", None, None, 40, 44])
    wb.save(topline_all_path)
    monkeypatch.setattr(appmod, "TOPLINE_FILES", {
        "all": topline_all_path, "md_est": topline_dir / "Category Topline - MD Estimates.xlsx",
    })
    appmod._REPORT_CACHE["topline_platforms"].clear()
    appmod._REPORT_CACHE["topline_rows"].clear()

    # Brand Share Cons: Cat A has 2 brands (BrandX + Others plug), Cat B (flat L0) has only an
    # Others row (the whole category is one unbranded bucket) - both groups have a locked cell.
    bsc_dir = portal_dir / "Brandshare Cons"
    bsc_dir.mkdir(parents=True)
    bsc_all_path = bsc_dir / "Brand Share Cons - All Platforms.xlsx"
    wb2 = openpyxl.Workbook()
    wb2.remove(wb2.active)
    ws2 = wb2.create_sheet("Blinkit")
    ws2.append(["Kind", "Level", "Label", "L0", "L1", "L2"] + months)
    ws2.append(["total", 0, "Total BPC", None, None, None, 100, 110])
    ws2.append(["l0", 1, "Cat A", "Cat A", None, None, 60, 66])
    ws2.append(["brand", 2, "BrandX", "Cat A", "Cat A", "Cat A", 42, 46.2])
    ws2.append(["others", 2, "Others", "Cat A", "Cat A", "Cat A", "=G3-G4", "=H3-H4"])
    ws2.append(["l0_flat", 1, "Cat B", "Cat B", None, None, 40, 44])
    ws2.append(["others", 2, "Others", "Cat B", "Cat B", "Cat B", "=G6", "=H6"])
    wb2.save(bsc_all_path)
    monkeypatch.setattr(appmod, "BSC_FILES", {
        "all": bsc_all_path, "md_est": bsc_dir / "Brand Share Cons - MD Estimates.xlsx",
    })
    appmod._REPORT_CACHE["bsc_platforms"].clear()
    appmod._REPORT_CACHE["bsc_rows"].clear()

    fake_md = pd.DataFrame({"_MonthLabel": ["Oct-30", "Nov-30"]})
    monkeypatch.setattr(appmod, "_get_master_data_df", lambda: fake_md)

    # "Expert data" (GMV/MoM/YoY benchmark + Brand Validation's Expert column) now comes from the
    # Raw expert data folder, not Master Data - reset per test to empty; tests that need expert
    # data set _CLIENT_RAW_AGG["df"] themselves.
    _CLIENT_RAW_AGG["df"] = pd.DataFrame(columns=["month_label", "l0", "l1", "l2", "brand", "gmv"])
    monkeypatch.setattr(appmod, "_load_raw_expert_agg", lambda platform: _CLIENT_RAW_AGG["df"])

    appmod.JOBS.clear()
    return TestClient(appmod.app)


_CLIENT_RAW_AGG = {"df": None}


def _lock(client, month, platforms, password="sl-pw"):
    return client.post(f"/monthly-setup/{month}/lock", json={"platforms": platforms, "password": password})


def test_set_total_gmv_propagates_category_and_computes_brand_rows_in_bsc(client):
    _lock(client, "Dec'30", {"Blinkit": False})
    r = client.post("/monthly-setup/Dec'30/Blinkit/category-topline/set-total-gmv", json={"total_gmv": 200})
    assert r.status_code == 200, r.text

    data = client.get("/brand-share-cons", params={"variant": "all", "platform": "Blinkit"}).json()
    assert "Dec-30" in data["months"]
    pos = data["months"].index("Dec-30")
    rows_by_label = {row["label"]: row for row in data["rows"]}

    # Step 1: Category Topline's own numbers (flat 0.6/0.4 history -> unchanged) landed in BSC.
    assert rows_by_label["Total BPC"]["gmv"][pos] == pytest.approx(200.0)
    assert rows_by_label["Cat A"]["gmv"][pos] == pytest.approx(120.0, abs=1e-4)  # 200 * 0.6
    assert rows_by_label["Cat B"]["gmv"][pos] == pytest.approx(80.0, abs=1e-4)

    # Step 2: brand rows under Cat A reconciled to Cat A's own just-written GMV (120).
    brand_x = rows_by_label["BrandX"]
    cat_a_others = [r for r in data["rows"] if r["label"] == "Others" and r["l0"] == "Cat A"][0]
    assert brand_x["share"][pos] == pytest.approx(0.7, abs=1e-6)  # flat 42/60 history
    assert brand_x["gmv"][pos] == pytest.approx(84.0, abs=1e-4)  # 120 * 0.7
    assert cat_a_others["gmv"][pos] == pytest.approx(36.0, abs=1e-4)
    assert brand_x["gmv"][pos] + cat_a_others["gmv"][pos] == pytest.approx(120.0, abs=1e-6)

    # Cat B's sole brand-level row is itself the plug (100% share) - reconciles to Cat B's GMV (80).
    cat_b_others = [r for r in data["rows"] if r["label"] == "Others" and r["l0"] == "Cat B"][0]
    assert cat_b_others["gmv"][pos] == pytest.approx(80.0, abs=1e-4)

    # The plug cells are real Excel formulas, not plain numbers (Manual Change must keep working).
    wb = openpyxl.load_workbook(appmod.BSC_FILES["all"])
    ws = wb["Blinkit"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = header.index("GMV Dec-30") + 1
    assert str(ws.cell(row=5, column=col_idx).value).startswith("=")  # Cat A's Others row
    assert str(ws.cell(row=7, column=col_idx).value).startswith("=")  # Cat B's Others row
    wb.close()


# ---------------------------------------------------------------- Brand Validation (step 3)

def test_brand_validation_rolls_up_bsc_and_excludes_others(client):
    _lock(client, "Dec'30", {"Blinkit": False})
    client.post("/monthly-setup/Dec'30/Blinkit/category-topline/set-total-gmv", json={"total_gmv": 200})

    r = client.get("/brand-validation", params={"platform": "Blinkit"})
    assert r.status_code == 200
    data = r.json()
    brands = {row["brand"]: row for row in data["rows"]}

    # "Others" is a generic per-category leftover, not a trackable brand - excluded entirely.
    assert "Others" not in brands
    assert "BrandX" in brands
    pos = data["months"].index("Dec-30")
    assert brands["BrandX"]["bsc_absolute"][pos] == pytest.approx(84.0, abs=1e-4)  # 120 * 0.7
    # No Master Data wired in this fixture at all for Blinkit - expert side is NA (None) throughout.
    assert brands["BrandX"]["expert_absolute"][pos] is None
    assert data["has_expert_data"] is False


def test_brand_validation_classifies_lo_tier_from_master_data_and_sorts_first(client, monkeypatch):
    # Master Data says "BrandX" (the fixture's own real BSC brand row) is an LO brand - the tier
    # classification must come from this live lookup, not a hardcoded name.
    fake_md = pd.DataFrame({
        "Platform": ["Blinkit"], "Brand": ["BrandX"], "LO Brand": ["LO"],
        "_MonthLabel": ["Oct-30"], "GMV (INR Cr.)": [10.0],
    })
    monkeypatch.setattr(appmod, "_get_master_data_df", lambda: fake_md)
    # Expert GMV itself now comes from the Raw expert data folder, not Master Data.
    _CLIENT_RAW_AGG["df"] = pd.DataFrame({
        "month_label": ["Oct-30"], "l0": ["cat a"], "l1": ["cat a"], "l2": ["cat a"],
        "brand": ["brandx"], "gmv": [10.0],
    })

    _lock(client, "Dec'31", {"Blinkit": False})
    client.post("/monthly-setup/Dec'31/Blinkit/category-topline/set-total-gmv", json={"total_gmv": 200})

    r = client.get("/brand-validation", params={"platform": "Blinkit"})
    assert r.status_code == 200
    data = r.json()
    assert data["has_expert_data"] is True
    brands = {row["brand"]: row for row in data["rows"]}
    assert brands["BrandX"]["tier"] == "lo"
    assert data["rows"][0]["brand"] == "BrandX"  # lo tier sorts first

    pos = data["months"].index("Oct-30")
    assert brands["BrandX"]["expert_absolute"][pos] == pytest.approx(10.0)


def test_brand_validation_merges_case_variant_brand_labels(client, monkeypatch):
    # Real production bug: the same brand can appear across several L2 rows with inconsistent
    # casing (e.g. "LANEIGE" in most categories, "Laneige" in one more) - these must be summed as
    # ONE brand, not fragmented into two rows that each then look like they're missing most of
    # their own GMV against Master Data's (correctly case-insensitive) total.
    wb = openpyxl.load_workbook(appmod.BSC_FILES["all"])
    ws = wb["Blinkit"]
    ws.append(["brand", 2, "ACME", "Cat A", "Cat A", "Cat A", 10, 11])
    ws.append(["brand", 2, "Acme", "Cat A", "Cat A", "Cat A", 5, 5.5])
    wb.save(appmod.BSC_FILES["all"])
    appmod._REPORT_CACHE["bsc_platforms"].clear()
    appmod._REPORT_CACHE["bsc_rows"].clear()

    fake_md = pd.DataFrame({
        "Platform": ["Blinkit"], "Brand": ["acme"], "LO Brand": ["Non-LO"],
        "_MonthLabel": ["Oct-30"], "GMV (INR Cr.)": [10.0],
    })
    monkeypatch.setattr(appmod, "_get_master_data_df", lambda: fake_md)
    _CLIENT_RAW_AGG["df"] = pd.DataFrame({
        "month_label": ["Oct-30"], "l0": ["cat a"], "l1": ["cat a"], "l2": ["cat a"],
        "brand": ["acme"], "gmv": [10.0],
    })

    r = client.get("/brand-validation", params={"platform": "Blinkit"})
    assert r.status_code == 200
    data = r.json()
    acme_rows = [row for row in data["rows"] if row["brand"].lower() == "acme"]
    assert len(acme_rows) == 1  # NOT two separate "ACME"/"Acme" rows
    pos = data["months"].index("Oct-30")
    assert acme_rows[0]["bsc_absolute"][pos] == pytest.approx(15.0)  # 10 + 5, summed across casings
    assert acme_rows[0]["expert_absolute"][pos] == pytest.approx(10.0)


def test_brand_validation_shows_real_tm_brand_from_master_data(client, monkeypatch):
    # TM Brand is Master Data's own real parent-company field - shown directly, not guessed.
    fake_md = pd.DataFrame({
        "Platform": ["Blinkit"], "Brand": ["BrandX"], "LO Brand": ["Non-LO"], "TM Brand": ["HUL"],
        "_MonthLabel": ["Oct-30"], "GMV (INR Cr.)": [10.0],
    })
    monkeypatch.setattr(appmod, "_get_master_data_df", lambda: fake_md)

    _lock(client, "Dec'32", {"Blinkit": False})
    client.post("/monthly-setup/Dec'32/Blinkit/category-topline/set-total-gmv", json={"total_gmv": 200})

    r = client.get("/brand-validation", params={"platform": "Blinkit"})
    assert r.status_code == 200
    brands = {row["brand"]: row for row in r.json()["rows"]}
    assert brands["BrandX"]["tier"] == "other"  # Non-LO
    assert brands["BrandX"]["tm_brand"] == "HUL"
